"""
Etapa 3 - Abre a planilha de equalizacao (gerada pela etapa 2), adiciona as
colunas "Preco Medio" e "VPL" (TMA 10%) na aba "Equalização" e calcula os
valores para cada fornecedor usando formulas nativas do Excel (auditavel e
recalculavel caso o usuario altere um preco na planilha).

Premissas do calculo de VPL (documentadas também na propria planilha):
  - Fluxo de caixa anual (custo) = Preço do ano (R$/MWh) x Volume (MWm) x
    8.760 h/ano (conversão padrão de MWm para MWh usada no setor).
  - Cada fluxo anual é descontado a partir da Data Base (01/07/2026): o ano
    de 2027 é o período 1, 2028 o período 2, e assim por diante -> mesma
    convenção da função NPV do Excel (fluxos ao final de cada período).
  - TMA (taxa mínima de atratividade) = 10% a.a.
  - O VPL é calculado com a função nativa NPV() do Excel.
"""
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import criar_planilha_equalizacao

BASE_DIR = Path(__file__).parent

TMA = 0.10
HORAS_POR_ANO = 8760
ANO_BASE = 2026  # Data base dos preços: 01/07/2026

# A aba "Equalização" começa na linha 2 / coluna B (linha 1 e coluna A ficam
# em branco como margem).
LINHA_CABECALHO = 2
COL_INICIAL = 2

AZUL_ESCURO = "1F4E78"
CINZA = "F2F2F2"
BORDA_FINA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _estilizar_cabecalho(ws, col, linha=LINHA_CABECALHO):
    cel = ws.cell(row=linha, column=col)
    cel.font = Font(color="FFFFFF", bold=True)
    cel.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cel.border = BORDA_FINA


def main():
    arquivo = criar_planilha_equalizacao.caminho_planilha_saida()
    wb = openpyxl.load_workbook(arquivo)
    ws = wb["Equalização"]

    n_colunas = ws.max_column
    n_linhas = ws.max_row

    cabecalho = {ws.cell(row=LINHA_CABECALHO, column=c).value: c for c in range(1, n_colunas + 1)}
    col_volume = cabecalho["Volume (MWm)"]
    colunas_preco = {
        int(m.group(1)): c
        for nome, c in cabecalho.items()
        if (m := re.match(r"Preço (\d{4}) \(R\$/MWh\)", nome or ""))
    }
    if not colunas_preco:
        raise RuntimeError("Nenhuma coluna de preço anual encontrada na aba Equalização.")

    # --- Novas colunas: Preço Médio e VPL (logo após a última coluna) ----
    col_preco_medio = n_colunas + 1
    col_vpl = n_colunas + 2

    # --- Parâmetros de cálculo (bloco abaixo da tabela) -------------------
    linha_param = n_linhas + 3
    parametros = [
        ("Parâmetros do cálculo de VPL", None),
        ("TMA (a.a.)", TMA),
        ("Horas por ano (MWm→MWh)", HORAS_POR_ANO),
        ("Ano base (Data Base 01/07)", ANO_BASE),
    ]
    for i, (rotulo, valor) in enumerate(parametros):
        linha = linha_param + i
        cel_rotulo = ws.cell(row=linha, column=COL_INICIAL, value=rotulo)
        cel_rotulo.font = Font(bold=True)
        if valor is not None:
            cel = ws.cell(row=linha, column=COL_INICIAL + 1, value=valor)
            if rotulo.startswith("TMA"):
                cel.number_format = "0%"
    cel_tma = ws.cell(row=linha_param + 1, column=COL_INICIAL + 1).coordinate
    cel_horas = ws.cell(row=linha_param + 2, column=COL_INICIAL + 1).coordinate

    ws.cell(row=LINHA_CABECALHO, column=col_preco_medio, value="Preço Médio do Período (R$/MWh)")
    ws.cell(row=LINHA_CABECALHO, column=col_vpl, value="VPL do Contrato (R$) TMA 10%")
    _estilizar_cabecalho(ws, col_preco_medio)
    _estilizar_cabecalho(ws, col_vpl)

    anos_ordenados = sorted(colunas_preco)
    col_letra_preco = {ano: ws.cell(row=LINHA_CABECALHO, column=colunas_preco[ano]).column_letter for ano in anos_ordenados}
    col_letra_volume = ws.cell(row=LINHA_CABECALHO, column=col_volume).column_letter

    for linha in range(LINHA_CABECALHO + 1, n_linhas + 1):
        vol_ref = f"{col_letra_volume}{linha}"
        primeiro_ano_col = col_letra_preco[anos_ordenados[0]]
        ultimo_ano_col = col_letra_preco[anos_ordenados[-1]]

        # Preço Médio: média simples dos preços anuais informados na proposta
        formula_media = f"=AVERAGE({primeiro_ano_col}{linha}:{ultimo_ano_col}{linha})"
        cel_media = ws.cell(row=linha, column=col_preco_medio, value=formula_media)
        cel_media.number_format = "#,##0.00"

        # VPL: NPV nativo do Excel sobre os custos anuais (preço x volume x
        # horas/ano), na ordem cronológica -> 2027 é o período 1, etc.
        termos = [
            f"({col_letra_preco[ano]}{linha}*{vol_ref}*${cel_horas})" for ano in anos_ordenados
        ]
        formula_vpl = f"=NPV(${cel_tma}, {', '.join(termos)})"
        cel_vpl = ws.cell(row=linha, column=col_vpl, value=formula_vpl)
        cel_vpl.number_format = "#,##0.00"

        for c in (col_preco_medio, col_vpl):
            cel = ws.cell(row=linha, column=c)
            cel.border = BORDA_FINA
            cel.alignment = Alignment(vertical="center")
            if linha % 2 == 0:
                cel.fill = PatternFill("solid", fgColor=CINZA)

    ws.column_dimensions[ws.cell(row=LINHA_CABECALHO, column=col_preco_medio).column_letter].width = 20
    ws.column_dimensions[ws.cell(row=LINHA_CABECALHO, column=col_vpl).column_letter].width = 20

    wb.save(arquivo)
    print(f"Colunas 'Preço Médio' e 'VPL' adicionadas em: {arquivo}")
    print(f"TMA utilizada: {TMA:.0%} | Horas/ano: {HORAS_POR_ANO} | Ano base: {ANO_BASE}")


if __name__ == "__main__":
    main()

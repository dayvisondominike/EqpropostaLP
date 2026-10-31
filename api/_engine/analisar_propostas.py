"""
Etapa 5 - Analisa a aderencia de cada proposta ao produto solicitado pelo
cliente (aba "Cliente") e grava o "Rank da Equalização" na aba "Equalização"
(1 = melhor colocada), além de um resumo executivo na aba "Cliente".

O rank é definido por preço: menor VPL primeiro, usando o Preço Médio como
critério de desempate. Todas as propostas entram no mesmo rank, independente
do status de aderência (aderência é só informativa nas colunas ao lado).
As linhas da tabela já são gravadas nessa ordem pela etapa 2
(criar_planilha_equalizacao.py) - aqui só é preciso rotular a posição.

Criterios "duros" de aderencia (reprovam a proposta se nao atendidos):
  Fonte, Submercado, Período, Volume (igual ao solicitado) e Modulação -
  valores-alvo parametrizáveis (ver constantes FONTE_PERCENTUAL,
  SUBMERCADO_REGEX/LABEL, ANO_INICIO/FIM, MODULACAO_DESEJADA abaixo).
  Exceção: Carga/Livre ofertada quando o pedido é Flat é mais ampla e só
  vira observação (o inverso reprova).

Criterios "leves" (viram observação, mas não desclassificam):
  Sazonalização e Flexibilidade diferentes do solicitado - podem ser mais
  restritivas (ponto de atenção) ou mais amplas (vantagem) que o pedido.
"""
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import criar_planilha_equalizacao

BASE_DIR = Path(__file__).parent

TMA = 0.10
HORAS_POR_ANO = 8760
ANO_BASE = 2026
TOLERANCIA_VOLUME = 0.001

# Critérios de aderência - parametrizáveis pela sidebar "Parâmetros" da
# interface Streamlit (streamlit_app.py monkeypatcha estas constantes antes
# de chamar avaliar_aderencia, do mesmo jeito que faz com TMA/HORAS_POR_ANO/
# ANO_BASE acima). Os valores abaixo são os padrões usados também quando o
# pipeline roda por linha de comando (App.py).
FONTE_PERCENTUAL = 50               # % mínimo de energia incentivada exigido
SUBMERCADO_REGEX = r"nordeste|\bNE\b"
SUBMERCADO_LABEL = "Nordeste"
ANO_INICIO = 2027
ANO_FIM = 2029
MODULACAO_DESEJADA = "Flat"
FLEXIBILIDADE_SUPERIOR_MINIMA = 30.0  # % mínima aceitável no limite superior (+)
FLEXIBILIDADE_INFERIOR_MINIMA = 30.0  # % mínima aceitável no limite inferior (-)
# Sazonalização "Flat" = sem tolerância (mínimos 0/0); "Personalizada" define
# os mínimos abaixo. Se a proposta oferecer MAIS tolerância que o mínimo
# exigido, isso não desclassifica nem gera ressalva - só uma nota informativa.
SAZONALIZACAO_SUPERIOR_MINIMA = 0.0   # % mínima exigida no limite superior (+)
SAZONALIZACAO_INFERIOR_MINIMA = 0.0   # % mínima exigida no limite inferior (-)

# A aba "Equalização" começa na linha 2 / coluna B (ver criar_planilha_equalizacao.py).
LINHA_CABECALHO = 2
COL_INICIAL = 2

AZUL_ESCURO = "1F4E78"
VERDE = "C6EFCE"
VERDE_FONTE = "006100"
VERMELHO = "FFC7CE"
VERMELHO_FONTE = "9C0006"
AMARELO = "FFEB9C"
AMARELO_FONTE = "9C6500"
VERDE_RANK1 = "2E7D32"
CINZA = "F2F2F2"
BORDA_FINA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _num(valor):
    if valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    return float(str(valor).replace(",", "."))


def _txt(valor) -> str:
    """Converte um campo de texto para string vazia com segurança. "valor or
    ''" não cobre NaN do pandas (colunas vazias em linhas "SEM DADOS" sem
    parser, como a Ludcom) porque NaN é "truthy" em Python - sobraria como
    float e quebraria nas chamadas de re.search/.strip() logo depois."""
    if valor is None or valor != valor:
        return ""
    return str(valor)


def ler_cliente(ws_cliente):
    dados = {}
    for linha in range(2, ws_cliente.max_row + 1):
        rotulo = ws_cliente.cell(row=linha, column=1).value
        valor = ws_cliente.cell(row=linha, column=2).value
        if rotulo:
            dados[rotulo] = valor
    return dados


def calcular_preco_medio_vpl(precos_por_ano: dict, volume: float):
    # Nem toda cotação tem os mesmos anos para todos os fornecedores (ex.:
    # a maioria manda proposta para 2027-2029, mas a Serena só manda para
    # 2027) - ignoramos os anos sem preço em vez de estourar ao somar None.
    anos = sorted(a for a, p in precos_por_ano.items() if p is not None)
    if not anos or volume is None:
        return None, None
    precos = [precos_por_ano[a] for a in anos]
    media = sum(precos) / len(precos)
    vpl = sum(
        p * volume * HORAS_POR_ANO / (1 + TMA) ** (a - ANO_BASE)
        for a, p in zip(anos, precos)
    )
    return media, vpl


def _parse_intervalo_percentual(texto: str):
    """Extrai os limites inferior (-) e superior (+) de um campo percentual
    bilateral da proposta (usado tanto para Flexibilidade quanto para
    Sazonalização), como magnitudes positivas. Cobre tanto notação simétrica
    ("+/-30%", "+-30%", "-/+ 30%", "+ou- 30%", "Variação de 30,00%") quanto
    assimétrica com sinais explícitos ("-100% / +100%", "LI: -50% LS: +30%").
    Retorna (None, None) se nada for encontrado (ex.: texto "Flat")."""
    texto = _txt(texto)

    combinado = re.search(
        r"(?:\+\s*/?\s*-|-\s*/?\s*\+|\+\s*ou\s*-|-\s*ou\s*\+)\s*(\d+(?:,\d+)?)\s*%", texto, re.IGNORECASE
    )
    if combinado:
        valor = float(combinado.group(1).replace(",", "."))
        return valor, valor

    # Faixa em % do montante em torno de 100% (ex.: CPFL: "90% a 110%" = ±10%).
    # Precisa vir antes dos pares com sinal: em "90% - 110%" o hífen é
    # separador da faixa, não o sinal de "-110%".
    faixa = re.search(r"(\d+(?:,\d+)?)\s*%\s*(?:a|até|-|–)\s*(\d+(?:,\d+)?)\s*%", texto, re.IGNORECASE)
    if faixa:
        de, ate = sorted(float(v.replace(",", ".")) for v in faixa.groups())
        if de <= 100 <= ate:
            return 100 - de, ate - 100

    pares = re.findall(r"([+\-])\s*(\d+(?:,\d+)?)\s*%", texto)
    if pares:
        positivos = [float(v.replace(",", ".")) for s, v in pares if s == "+"]
        negativos = [float(v.replace(",", ".")) for s, v in pares if s == "-"]
        superior = max(positivos) if positivos else None
        inferior = max(negativos) if negativos else None
        if superior is None:
            superior = inferior
        if inferior is None:
            inferior = superior
        if superior is not None and inferior is not None:
            return inferior, superior

    # Convenção "X% da energia mensal sazonalizada" / take-or-pay (ex.: CEMIG:
    # "130% da energia mensal sazonalizada (take-or-pay 60%)", "Take-or-pay 0% /
    # Flexibilidade mensal 130%") e "Mínimo de X% e Máximo de Y%" (ex.:
    # Auren): os valores são percentuais do montante contratado (100% = sem
    # flexibilidade) - abaixo de 100% é o piso de consumo (flexibilidade
    # inferior = 100% menos o piso; o take-or-pay é esse piso) e acima de 100%
    # é o teto (flexibilidade superior = teto menos 100%).
    valores_absolutos = [
        float(v.replace(",", "."))
        for padrao in (
            r"(\d+(?:,\d+)?)\s*%\s*(?:da|sobre\s+(?:a\s+)?)\s*energia\s+mensal",
            r"take[\s-]*or[\s-]*pay\s*:?\s*(?:de\s*)?(\d+(?:,\d+)?)\s*%",
            r"flexibilidade\s+mensal\s*:?\s*(?:de\s*)?(\d+(?:,\d+)?)\s*%",
        )
        for v in re.findall(padrao, texto, re.IGNORECASE)
    ]
    minimo = re.search(r"M[íi]nimo\s+(?:de\s*)?(\d+(?:,\d+)?)\s*%", texto, re.IGNORECASE)
    maximo = re.search(r"M[áa]ximo\s+(?:de\s*)?(\d+(?:,\d+)?)\s*%", texto, re.IGNORECASE)
    if minimo:
        valores_absolutos.append(float(minimo.group(1).replace(",", ".")))
    if maximo:
        valores_absolutos.append(float(maximo.group(1).replace(",", ".")))
    if valores_absolutos:
        abaixo = [v for v in valores_absolutos if v < 100]
        acima = [v for v in valores_absolutos if v > 100]
        inferior = (100 - min(abaixo)) if abaixo else None
        superior = (max(acima) - 100) if acima else None
        if superior is None:
            superior = inferior
        if inferior is None:
            inferior = superior
        if inferior is not None and superior is not None:
            return inferior, superior

    valores = re.findall(r"(\d+(?:,\d+)?)\s*%", texto)
    if valores:
        valor = max(float(v.replace(",", ".")) for v in valores)
        return valor, valor

    return None, None


def _normalizar_modulacao(texto: str) -> str:
    """Fornecedores nomeiam a modulação por curva de carga de jeitos diferentes
    (ex.: CEMIG: "Pela Curva de Carga – Modulação Livre", que a extração reduz
    a "Livre") - "Livre" e qualquer menção a "carga" valem como "Carga" na
    comparação de aderência."""
    texto = _txt(texto).strip()
    if re.search(r"carga|\blivre\b", texto, re.IGNORECASE):
        return "Carga"
    if re.search(r"\bflat\b", texto, re.IGNORECASE):
        return "Flat"
    return texto


def avaliar_aderencia(row: dict, cliente: dict):
    motivos_reprovacao = []
    observacoes = []    # ressalvas: rebaixam o status para "Aderente (com ressalvas)"
    informativas = []   # não rebaixam o status - só entram na observação final

    fonte = _txt(row["Fonte"])
    if str(FONTE_PERCENTUAL) not in fonte:
        motivos_reprovacao.append(f"Fonte '{fonte}' diferente da solicitada (Incentivada {FONTE_PERCENTUAL}%)")

    submercado = _txt(row["Submercado"])
    if not re.search(SUBMERCADO_REGEX, submercado, re.IGNORECASE):
        motivos_reprovacao.append(f"Submercado '{submercado}' diferente do solicitado ({SUBMERCADO_LABEL})")

    periodo = _txt(row["Período"])
    if not (re.search(str(ANO_INICIO), periodo) and re.search(str(ANO_FIM), periodo)):
        motivos_reprovacao.append(f"Período '{periodo}' não cobre integralmente {ANO_INICIO}-{ANO_FIM}")

    volume_cliente = _num(cliente.get("Volume (MWm)"))
    volume_proposta = _num(row["Volume (MWm)"])
    if row.get("Volume informado na proposta") in (False, "False"):
        informativas.append(f"Volume não informado na proposta - considerado o volume solicitado ({volume_cliente} MWm)")
    if volume_proposta is None or volume_cliente is None or abs(volume_proposta - volume_cliente) > TOLERANCIA_VOLUME:
        motivos_reprovacao.append(
            f"Volume {volume_proposta} MWm diferente do solicitado ({volume_cliente} MWm)"
        )

    modulacao_original = _txt(row["Modulação"])
    modulacao = _normalizar_modulacao(modulacao_original)
    desejada = MODULACAO_DESEJADA.strip()
    if modulacao.lower() == desejada.lower():
        pass
    elif modulacao == "Carga" and desejada.lower() == "flat":
        # Modulação pela curva de carga cobre também um consumo flat.
        informativas.append(f"Modulação ofertada ({modulacao_original}) é mais ampla que a solicitada ({desejada})")
    else:
        motivos_reprovacao.append(f"Modulação '{modulacao_original}' diferente da solicitada ({desejada})")

    sazo = _txt(row["Sazonalidade"])
    sazo_inferior, sazo_superior = _parse_intervalo_percentual(sazo)
    sazo_superior = sazo_superior if sazo_superior is not None else 0.0
    sazo_inferior = sazo_inferior if sazo_inferior is not None else 0.0

    sazo_restritiva = []
    sazo_ampla = []
    if sazo_superior < SAZONALIZACAO_SUPERIOR_MINIMA:
        sazo_restritiva.append(f"limite superior ofertado +{sazo_superior:g}% é menor que o mínimo solicitado (+{SAZONALIZACAO_SUPERIOR_MINIMA:g}%)")
    elif sazo_superior > SAZONALIZACAO_SUPERIOR_MINIMA:
        sazo_ampla.append(f"limite superior ofertado +{sazo_superior:g}% é maior que o mínimo solicitado (+{SAZONALIZACAO_SUPERIOR_MINIMA:g}%)")
    if sazo_inferior < SAZONALIZACAO_INFERIOR_MINIMA:
        sazo_restritiva.append(f"limite inferior ofertado -{sazo_inferior:g}% é menor que o mínimo solicitado (-{SAZONALIZACAO_INFERIOR_MINIMA:g}%)")
    elif sazo_inferior > SAZONALIZACAO_INFERIOR_MINIMA:
        sazo_ampla.append(f"limite inferior ofertado -{sazo_inferior:g}% é maior que o mínimo solicitado (-{SAZONALIZACAO_INFERIOR_MINIMA:g}%)")

    if sazo_restritiva:
        observacoes.append(f"Sazonalização ofertada ({sazo}) é mais restritiva que a solicitada: {'; '.join(sazo_restritiva)}")
    if sazo_ampla:
        informativas.append(f"Sazonalização ofertada ({sazo}) é mais ampla que a solicitada: {'; '.join(sazo_ampla)}")

    flex = _txt(row["Flexibilidade"])
    flex_inferior, flex_superior = _parse_intervalo_percentual(flex)
    if flex_inferior is None and flex_superior is None and re.search(r"\bflat\b|sem\s+flexibilidade", flex, re.IGNORECASE):
        # Flexibilidade "Flat" equivale a +-0% (nenhuma tolerância de consumo).
        flex_inferior = flex_superior = 0.0
    flex_restritiva = []
    flex_ampla = []
    if flex_superior is not None:
        if flex_superior < FLEXIBILIDADE_SUPERIOR_MINIMA:
            flex_restritiva.append(f"limite superior +{flex_superior:g}% menor que o solicitado (+{FLEXIBILIDADE_SUPERIOR_MINIMA:g}%)")
        elif flex_superior > FLEXIBILIDADE_SUPERIOR_MINIMA:
            flex_ampla.append(f"limite superior ofertado +{flex_superior:g}% é maior que o solicitado (+{FLEXIBILIDADE_SUPERIOR_MINIMA:g}%)")
    if flex_inferior is not None:
        if flex_inferior < FLEXIBILIDADE_INFERIOR_MINIMA:
            flex_restritiva.append(f"limite inferior -{flex_inferior:g}% menor que o solicitado (-{FLEXIBILIDADE_INFERIOR_MINIMA:g}%)")
        elif flex_inferior > FLEXIBILIDADE_INFERIOR_MINIMA:
            flex_ampla.append(f"limite inferior ofertado -{flex_inferior:g}% é maior que o solicitado (-{FLEXIBILIDADE_INFERIOR_MINIMA:g}%)")
    if flex_restritiva:
        observacoes.append(f"Flexibilidade ofertada ({flex}) é mais restritiva que a solicitada: {'; '.join(flex_restritiva)}")
    if flex_ampla:
        informativas.append(f"Flexibilidade ofertada ({flex}) é mais ampla que a solicitada: {'; '.join(flex_ampla)}")

    if motivos_reprovacao:
        status = "Não aderente"
    elif observacoes:
        status = "Aderente (com ressalvas)"
    else:
        status = "Aderente"

    todas_notas = motivos_reprovacao + observacoes + informativas
    obs_final = "; ".join(todas_notas) if todas_notas else "Sem divergências relevantes"
    return status, obs_final


def _colorir_aderencia(cel, status):
    if status.startswith("Aderente") and "ressalvas" not in status:
        cel.fill = PatternFill("solid", fgColor=VERDE)
        cel.font = Font(color=VERDE_FONTE)
    elif "ressalvas" in status:
        cel.fill = PatternFill("solid", fgColor=AMARELO)
        cel.font = Font(color=AMARELO_FONTE)
    elif status.startswith("Não"):
        cel.fill = PatternFill("solid", fgColor=VERMELHO)
        cel.font = Font(color=VERMELHO_FONTE)


def main():
    arquivo = criar_planilha_equalizacao.caminho_planilha_saida()
    wb = openpyxl.load_workbook(arquivo)
    ws = wb["Equalização"]
    ws_cliente = wb["Cliente"]
    cliente = ler_cliente(ws_cliente)

    cabecalho = {ws.cell(row=LINHA_CABECALHO, column=c).value: c for c in range(1, ws.max_column + 1)}

    # A aba tem um bloco de parâmetros (TMA, horas/ano, ano base) abaixo da
    # tabela de propostas (escrito por calcular_equalizacao.py) - aqui
    # contamos só as linhas contíguas de fornecedores, ignorando esse bloco.
    n_linhas = LINHA_CABECALHO
    while ws.cell(row=n_linhas + 1, column=COL_INICIAL).value not in (None, ""):
        n_linhas += 1

    colunas_preco = {
        int(m.group(1)): cabecalho[nome]
        for nome in cabecalho
        if (m := re.match(r"Preço (\d{4}) \(R\$/MWh\)", nome or ""))
    }

    col_aderencia = max(cabecalho.values()) + 1
    col_obs = col_aderencia + 1
    col_rank = col_aderencia + 2

    for col, titulo in (
        (col_aderencia, "Aderente ao Produto do Cliente"),
        (col_obs, "Observações de Aderência"),
        (col_rank, "Rank da Equalização"),
    ):
        cel = ws.cell(row=LINHA_CABECALHO, column=col, value=titulo)
        cel.font = Font(color="FFFFFF", bold=True)
        cel.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = BORDA_FINA

    # As linhas já vêm ordenadas por VPL crescente (desempate por Preço Médio)
    # desde a etapa 2 - aqui só avaliamos aderência e numeramos a posição.
    propostas = []  # (linha, fornecedor, preco_medio, vpl)
    for linha in range(LINHA_CABECALHO + 1, n_linhas + 1):
        row = {nome: ws.cell(row=linha, column=c).value for nome, c in cabecalho.items()}
        if row.get("Volume (MWm)") in (None, ""):
            status, obs = "Sem dados", "Arquivo de proposta corrompido/vazio - fornecedor não pôde ser avaliado"
            media = vpl = None
        else:
            status, obs = avaliar_aderencia(row, cliente)
            precos_ano = {a: _num(row[ws.cell(row=LINHA_CABECALHO, column=c).value]) for a, c in colunas_preco.items()}
            media, vpl = calcular_preco_medio_vpl(precos_ano, _num(row["Volume (MWm)"]))
            propostas.append((linha, row["Fornecedor"], media, vpl))

        cel_ad = ws.cell(row=linha, column=col_aderencia, value=status)
        cel_ad.border = BORDA_FINA
        cel_ad.alignment = Alignment(horizontal="center", vertical="center")
        _colorir_aderencia(cel_ad, status)

        cel_obs = ws.cell(row=linha, column=col_obs, value=obs)
        cel_obs.border = BORDA_FINA
        cel_obs.alignment = Alignment(vertical="center", wrap_text=True)

        cel_rank = ws.cell(row=linha, column=col_rank)
        cel_rank.border = BORDA_FINA
        cel_rank.alignment = Alignment(horizontal="center", vertical="center")

        if linha % 2 == 0:
            for c in (col_aderencia, col_obs, col_rank):
                if ws.cell(row=linha, column=c).fill.fgColor.rgb in (None, "00000000"):
                    ws.cell(row=linha, column=c).fill = PatternFill("solid", fgColor=CINZA)

    ws.column_dimensions[ws.cell(row=LINHA_CABECALHO, column=col_aderencia).column_letter].width = 22
    ws.column_dimensions[ws.cell(row=LINHA_CABECALHO, column=col_obs).column_letter].width = 45
    ws.column_dimensions[ws.cell(row=LINHA_CABECALHO, column=col_rank).column_letter].width = 14

    # Numera o rank pela posição na tabela (já ordenada por VPL na etapa 2).
    for posicao, (linha, _forn, _media, _vpl) in enumerate(propostas, start=1):
        cel_rank = ws.cell(row=linha, column=col_rank, value=posicao)
        if posicao == 1:
            cel_rank.font = Font(bold=True, color="FFFFFF")
            cel_rank.fill = PatternFill("solid", fgColor=VERDE_RANK1)
    for linha in range(LINHA_CABECALHO + 1, n_linhas + 1):
        if ws.cell(row=linha, column=col_rank).value is None:
            ws.cell(row=linha, column=col_rank, value="-")

    # --- Resumo executivo na aba Cliente ---------------------------------
    linha_resumo = ws_cliente.max_row + 2
    ws_cliente.cell(row=linha_resumo, column=1, value="RESULTADO DA EQUALIZAÇÃO").font = Font(bold=True, size=12, color="FFFFFF")
    ws_cliente.cell(row=linha_resumo, column=1).fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    ws_cliente.merge_cells(start_row=linha_resumo, start_column=1, end_row=linha_resumo, end_column=2)

    if propostas:
        _linha1, forn1, media1, vpl1 = propostas[0]
        resumo = [
            ("1º colocado no Rank", forn1),
            ("Preço médio do período", f"R$ {media1:,.2f} /MWh"),
            ("VPL do contrato (TMA 10%)", f"R$ {vpl1:,.2f}"),
            (
                "Critério de decisão",
                "Rank por menor VPL (TMA 10%), com desempate pelo Preço Médio; todas as "
                "propostas recebidas entram no rank, independente da aderência ao produto "
                "solicitado (a aderência fica registrada nas colunas ao lado)",
            ),
            (
                "Ranking completo (1º ao último)",
                ", ".join(f"{i}. {f} (VPL R$ {v:,.2f})" for i, (_l, f, _m, v) in enumerate(propostas, start=1)),
            ),
        ]
    else:
        resumo = [("Resultado", "Nenhuma proposta com dados válidos foi encontrada.")]

    for i, (rotulo, valor) in enumerate(resumo, start=1):
        ws_cliente.cell(row=linha_resumo + i, column=1, value=rotulo).font = Font(bold=True)
        ws_cliente.cell(row=linha_resumo + i, column=2, value=valor).alignment = Alignment(wrap_text=True)

    wb.save(arquivo)
    if propostas:
        print(f"Rank da Equalização gravado. 1º colocado: {propostas[0][1]} (VPL R$ {propostas[0][3]:,.2f})")
    else:
        print("Nenhuma proposta com dados válidos para ranquear.")


if __name__ == "__main__":
    main()

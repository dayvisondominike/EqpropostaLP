"""
Leitura determinística (sem LLM) de planilhas de "Cotações Semanais": uma única
planilha com várias linhas, cada uma uma cotação de mercado de um fornecedor
(ex.: "CEMIG - Opção 2") com os preços por ano nas colunas 2026, 2027...,
Flexibilidade, Sazonalidade, Modulação e Fonte de Energia.

Cada linha vira uma proposta no mesmo schema que a extração dos PDFs produz.
Quando o mesmo fornecedor/opção aparece em mais de uma data (histórico
semanal), vale a cotação mais recente.
"""
import csv
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Optional

TIPO_COTACAO_MERCADO = "Cotação de mercado"


def _sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def _chave(texto) -> str:
    return re.sub(r"[^a-z0-9]", "", _sem_acento(str(texto or "")).lower())


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _numero(valor) -> Optional[float]:
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return float(valor)
    t = _texto(valor)
    if not t or t in ("-", "–", "—"):
        return None
    t = t.replace("R$", "").replace(" ", "")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".")
    else:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def _data(valor) -> Optional[date]:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    m = re.match(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})", _texto(valor))
    if m:
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    m = re.match(r"^\s*(\d{4})-(\d{2})-(\d{2})", _texto(valor))
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def _percentual_bilateral(valor) -> Optional[str]:
    """'30% / -50%' -> '+30% / -50%' (formato que a análise de aderência entende: limite superior / inferior)."""
    t = _texto(valor)
    if not t:
        return None
    m = re.match(r"^\s*([+-]?\d+(?:[.,]\d+)?)\s*%\s*/\s*([+-]?\d+(?:[.,]\d+)?)\s*%\s*$", t)
    if m:
        sup = abs(float(m.group(1).replace(",", ".")))
        inf = abs(float(m.group(2).replace(",", ".")))
        return f"+{sup:g}% / -{inf:g}%"
    return "Flat" if t.lower() == "flat" else t


def _modulacao(valor) -> Optional[str]:
    t = _texto(valor)
    if not t:
        return None
    if t.lower() == "flat":
        return "Flat"
    if t.lower() == "declarada":
        return "Carga (curva declarada)"
    return t.capitalize()


def _linhas(caminho: Path) -> list:
    if caminho.suffix.lower() == ".csv":
        bruto = caminho.read_bytes()
        for cod in ("utf-8-sig", "cp1252"):
            try:
                texto = bruto.decode(cod)
                break
            except UnicodeDecodeError:
                continue
        else:
            texto = bruto.decode("latin-1", errors="replace")
        delim = ";" if texto.count(";") >= texto.count(",") else ","
        return [list(l) for l in csv.reader(texto.splitlines(), delimiter=delim)]
    from openpyxl import load_workbook
    wb = load_workbook(caminho, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            linhas = [list(r) for r in ws.iter_rows(values_only=True)]
            if _achar_cabecalho(linhas) is not None:
                return linhas
        return []
    finally:
        wb.close()


def _achar_cabecalho(linhas: list) -> Optional[int]:
    for i, linha in enumerate(linhas[:30]):
        chaves = [_chave(c) for c in linha]
        tem_anos = sum(1 for c in linha if re.fullmatch(r"20\d{2}", _texto(c).split(".")[0]))
        if "fornecedores" in chaves or "fornecedor" in chaves:
            if tem_anos >= 1:
                return i
    return None


def ler_tabela_cotacoes(caminho) -> list:
    """Propostas (uma por fornecedor/opção, a cotação mais recente) ou [] se o arquivo não tiver esse formato."""
    caminho = Path(caminho)
    try:
        linhas = _linhas(caminho)
    except Exception:  # noqa: BLE001  (arquivo ilegível: segue o fluxo normal da LLM)
        return []
    topo = _achar_cabecalho(linhas)
    if topo is None:
        return []

    cab = linhas[topo]
    col = {}
    anos = {}
    for j, c in enumerate(cab):
        t = _texto(c)
        if re.fullmatch(r"20\d{2}(\.0)?", t):
            anos[int(float(t))] = j
        else:
            col.setdefault(_chave(t), j)

    def celula(linha, nome):
        j = col.get(nome)
        return linha[j] if j is not None and j < len(linha) else None

    nome_col = "fornecedores" if "fornecedores" in col else "fornecedor"
    melhores: dict = {}
    for linha in linhas[topo + 1:]:
        nome = _texto(celula(linha, nome_col))
        if not nome:
            continue
        precos = {a: _numero(linha[j]) for a, j in anos.items() if j < len(linha)}
        precos = {a: p for a, p in precos.items() if p is not None}
        if not precos:
            continue
        quando = _data(celula(linha, "datacadastro"))
        atual = melhores.get(_chave(nome))
        if atual is None or (quando or date.min) >= (atual["quando"] or date.min):
            melhores[_chave(nome)] = {"nome": nome, "linha": linha, "precos": precos, "quando": quando}

    propostas = []
    for item in melhores.values():
        linha, precos = item["linha"], item["precos"]
        id_cotacao = _texto(celula(linha, "id"))
        fonte = _texto(celula(linha, "fontedeenergia")) or None
        proposta = {
            "Fornecedor": item["nome"].upper(),
            "Tipo de Proposta": f"{TIPO_COTACAO_MERCADO} (ID {id_cotacao})" if id_cotacao else TIPO_COTACAO_MERCADO,
            "Fonte": fonte,
            "Submercado": None,  # a cotação semanal não traz submercado: assume o do cliente (ver engine.consolidar)
            "Período": f"01/01/{min(precos)} a 31/12/{max(precos)}",
            "Volume (MWm)": None,
            "Preços por ano": {str(a): p for a, p in sorted(precos.items())},
            "Flexibilidade": _percentual_bilateral(celula(linha, "flexibilidade")),
            "Sazonalidade": _percentual_bilateral(celula(linha, "sazonalidade")),
            "Modulação": _modulacao(celula(linha, "modulacao")),
            "Garantia": None,
            "RETUSD (R$/MWh)": None,
            "Pagamento": None,
            "Data Base": None,
            "Data da Proposta": item["quando"].strftime("%d/%m/%Y") if item["quando"] else None,
            "Validade da Proposta": None,
            "Cotação de mercado": True,
        }
        propostas.append(proposta)
    return sorted(propostas, key=lambda p: p["Fornecedor"])
